from docx import Document
import pandas as pd
import os

#! Update the file path before running the script
template_path = '[insert-path-here]'
if not os.path.isfile(template_path):
    raise FileNotFoundError(f"The file at {template_path} does not exist.")

#! Update the file path before running the script
# Load Excel data into a DataFrame
excel_path = '[insert-path-here]'
if not os.path.isfile(excel_path):
    raise FileNotFoundError(f"The file at {excel_path} does not exist.")

df = pd.read_excel(excel_path)

#! Update the file path before running the script
# Output directory for merged documents
output_dir = '[insert-path-here]'
os.makedirs(output_dir, exist_ok=True)  # Create the output folder if it doesn't exist

# Perform mail merge
for index, row in df.iterrows():
    company = str(row['Company'])
    address = str(row['Address'])

    # Create a copy of the document for each iteration
    merged_doc = Document(template_path)

    # Replace placeholders in the Word document
    for paragraph in merged_doc.paragraphs:
        for run in paragraph.runs:
            if '<<Company>>' in run.text:
                run.text = run.text.replace('<<Company>>', company)
            if '<<Address>>' in run.text:
                run.text = run.text.replace('<<Address>>', address)

    # Save the merged document
    output_filename = f"{company}.docx"
    output_path = os.path.join(output_dir, output_filename)
    merged_doc.save(output_path)
